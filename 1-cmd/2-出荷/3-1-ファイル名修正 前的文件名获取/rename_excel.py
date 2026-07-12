from pathlib import Path
import openpyxl


# =========================
# 要搜索的多个文件夹
# =========================

folders = [
    Path(r"C:\c_wk\10_会社\PDF-相关\Test\out"),
    Path(r"C:\c_wk\10_会社\PDF-相关\Test\out - コピー")
]


# 输出Excel
excel_file = Path(
    r"C:\c_wk\10_会社\PDF-相关\PDF一覧.xlsx"
)


# =========================
# 创建Excel
# =========================

wb = openpyxl.Workbook()

ws = wb.active

ws.title = "PDF一覧"


# 标题

ws.append([
    "文件夹",
    "PDF文件名",
    "完整路径"
])


# =========================
# 读取PDF
# =========================

row = 2


for folder in folders:

    for pdf_file in folder.glob("*.pdf"):


        ws.cell(
            row=row,
            column=1,
            value=folder.name
        )


        ws.cell(
            row=row,
            column=2,
            value=pdf_file.name
        )


        ws.cell(
            row=row,
            column=3,
            value=str(pdf_file)
        )


        row += 1



# 保存

wb.save(excel_file)


print("完成")
print("数量:", row-2)
print(excel_file)

input("按 Enter 结束...")